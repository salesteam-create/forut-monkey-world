#!/usr/bin/env python3
"""Voice tooling for Monkis verden.

  python3 tools/voice.py manifest   # after adding recordings: rebuild assets/voice/manifest.json
  python3 tools/voice.py script     # rebuild the recording script sheet (voice/voice-script.xlsx)

Recordings go in assets/voice/no/ and assets/voice/en/, named as in the script sheet
(for example hubGreeting.mp3). .mp3, .m4a, .wav and .ogg all work.
"""
import json
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
VOICE = ROOT / 'assets' / 'voice'
AUDIO = {'.mp3', '.m4a', '.wav', '.ogg', '.aac'}
LANGS = ('no', 'en')


def load_lines():
    # i18n.js and voice-lines.js are plain browser scripts; let node evaluate them.
    js = """
      global.window = {};
      require('./js/i18n.js'); require('./js/voice-lines.js');
      const get = (o, k) => k.split('.').reduce((a, b) => (a == null ? a : a[b]), o);
      console.log(JSON.stringify(window.VOICE_LINES.map(l => Object.assign({}, l, {
        file: window.voiceFile(l.key),
        no: get(window.I18N.no, l.key),
        en: get(window.I18N.en, l.key)
      }))));
    """
    out = subprocess.run(['node', '-e', js], cwd=ROOT, capture_output=True, text=True, check=True)
    return json.loads(out.stdout)


def manifest():
    data = {}
    for lang in LANGS:
        files = sorted(p for p in (VOICE / lang).glob('*') if p.suffix.lower() in AUDIO)
        data[lang] = {p.stem: f'{lang}/{p.name}' for p in files}
    (VOICE / 'manifest.json').write_text(json.dumps(data, indent=2, ensure_ascii=False) + '\n')
    known = {l['file'] for l in load_lines()}
    for lang in LANGS:
        extra = sorted(set(data[lang]) - known)
        missing = sorted(known - set(data[lang]))
        print(f'{lang}: {len(data[lang])} recordings, {len(missing)} lines still use the browser voice')
        if extra:
            print(f'  not matching any line (check the file name): {", ".join(extra)}')


def script():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    lines = load_lines()
    wb = Workbook()
    ws = wb.active
    ws.title = 'Script'
    head = ['#', 'Demo', 'Character', 'Screen', 'File name', 'Norwegian line', 'English line', 'Tone', 'Recorded NO', 'Recorded EN']
    widths = [5, 8, 20, 22, 22, 55, 55, 34, 13, 13]
    base = Font(name='Arial', size=10)
    bold = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    ws.append(head)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        c = ws.cell(1, i)
        c.font = bold
        c.fill = PatternFill('solid', fgColor='1B62B3')
        c.alignment = Alignment(vertical='center', wrap_text=True)
    demo_fill = PatternFill('solid', fgColor='FFF4C7')
    for n, l in enumerate(lines, 1):
        ws.append([n, 'Yes' if l['demo'] else '', l['char'], l['screen'], l['file'] + '.mp3',
                   l['no'], l['en'], l['tone'], '', ''])
        for col in range(1, len(head) + 1):
            c = ws.cell(n + 1, col)
            c.font = base
            c.alignment = Alignment(vertical='top', wrap_text=True)
            if l['demo']:
                c.fill = demo_fill
    dv = DataValidation(type='list', formula1='"Yes,No"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f'I2:J{len(lines) + 1}')
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:J{len(lines) + 1}'

    g = wb.create_sheet('How to record')
    g.column_dimensions['A'].width = 110
    guide = [
        ('How to record Monki', True),
        ('Rows shaded yellow (Demo = Yes) are heard in the pitch walkthrough. Record those first.', False),
        ('Record in a quiet room with soft furnishings. A phone held 20 to 30 cm from the mouth is enough.', False),
        ('Read each line as if talking to a 4 year old sitting next to you: warm, clear and a little slower than normal.', False),
        ('Leave half a second of silence before and after each line, and do two takes of anything you stumble on.', False),
        ('Palaiya (breathing lines) should sound slower and calmer than Monki. The secret friend should sound different from Monki if possible.', False),
        ('Save each line as its own file named exactly as in the File name column, for example hubGreeting.mp3. .m4a from an iPhone voice memo is fine too.', False),
        ('Norwegian files go in assets/voice/no/, English files in assets/voice/en/. Then run: python3 tools/voice.py manifest', False),
        ('Any line without a recording still plays with the device voice, so partial recordings are fine.', False),
        ('Mark the Recorded NO / Recorded EN columns as you go.', False),
    ]
    for r, (text, is_title) in enumerate(guide, 1):
        c = g.cell(r, 1, text)
        c.font = Font(name='Arial', size=14 if is_title else 10, bold=is_title)
        c.alignment = Alignment(wrap_text=True, vertical='top')

    out = ROOT / 'voice' / 'voice-script.xlsx'
    out.parent.mkdir(exist_ok=True)
    wb.save(out)
    print(f'Wrote {out.relative_to(ROOT)} with {len(lines)} lines ({sum(l["demo"] for l in lines)} for the demo)')


if __name__ == '__main__':
    cmd = sys.argv[1] if len(sys.argv) > 1 else ''
    if cmd == 'manifest':
        manifest()
    elif cmd == 'script':
        script()
    else:
        print(__doc__)
